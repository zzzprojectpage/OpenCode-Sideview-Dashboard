"""Generate Instructions.docx for the GitHub release.

Reproducible: embeds the screenshots extracted from the author's own desktop
(build/docx-media/) with captions describing what each panel section shows.
Run:  py build_instructions.py   ->  writes Instructions.docx next to this script.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
MEDIA = ROOT / 'build/docx-media'
OUT = ROOT / 'Instructions.docx'

SCREENSHOTS = [
    ('image1.png',
     'Context, cost, speed, cache and MCP usage',
     'Context shows the percentage of the model context window used by the latest '
     'measured request, with token counts underneath. Cost · session is the cost OpenCode '
     'itself reports for the whole session. Speed is the last completed generation in '
     'tokens per second (K = thousand, M = million). Cache hit ratio is cached reads over '
     'all input tokens. MCPs lists every configured server and how many calls each one '
     'served in the active context.'),
    ('image2.png',
     'OpenCode Go caps with the estimated per-model split',
     'Each cap window shows the percentage used as reported by the provider, the '
     'percentage left, and the reset time. Underneath, BY MODEL · ESTIMATED lists which '
     'models consumed the window: the provider total is distributed by each Go model\u2019s '
     'share of locally recorded cost. Free models show 0% because they carry no billable '
     'cost. These shares are estimates, not provider data.'),
    ('image3.png',
     'OpenAI caps, model token share and diagnostics',
     'The OpenAI section works the same way: 5-hour, weekly and monthly windows with '
     'reset times (a window the provider does not report shows Unavailable instead of a '
     'made-up number). Model token share breaks the active context down by model. The '
     'small line at the bottom reports the detected session and how many client calls '
     'the sidebar has observed - useful when asking for help.'),
]


def heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = None
    return paragraph


def bullets(document, items):
    for item in items:
        document.add_paragraph(item, style='List Bullet')


def main():
    missing = [name for name, _, _ in SCREENSHOTS if not (MEDIA / name).is_file()]
    if missing:
        raise SystemExit(f'missing screenshots in {MEDIA}: {missing}')

    document = Document()
    style = document.styles['Normal']
    style.font.size = Pt(11)

    title = document.add_heading('OpenCode Session Telemetry Sidebar', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Installation and user guide')
    run.font.size = Pt(14)

    heading(document, 'What this is', level=1)
    document.add_paragraph(
        'A right-hand sidebar for the OpenCode desktop app. It shows live session '
        'telemetry: context-window usage, session cost, generation speed, cache hit '
        'ratio, MCP usage, your OpenCode Go and OpenAI subscription caps (5-hour, '
        'weekly, monthly), and an estimated per-model share of each Go cap window. '
        'A Search button beside it finds text across every past session.')
    bullets(document, [
        'Context window in percent and tokens, for the latest measured request.',
        'Session cost, taken from OpenCode\u2019s own session record.',
        'Speed of the last completed generation, in tokens/second.',
        'Cache hit ratio (cached reads over all input tokens).',
        'MCP servers with per-server call counts in the active context.',
        'OpenCode Go and OpenAI caps with reset times and remaining percent.',
        'Estimated per-model share of each Go cap window.',
        'History search over user prompts and agent replies, opening hits as new tabs.',
    ])

    heading(document, 'Requirements', level=1)
    bullets(document, [
        'Windows 10 or 11.',
        'The OpenCode desktop app, version 2.x.',
        'Python 3.8 or newer, available as py on PATH (check with: py --version).',
    ])

    heading(document, 'Before you start', level=1)
    document.add_paragraph(
        'Quit OpenCode completely: close every window and the tray icon. '
        'The installer refuses to run while OpenCode is open, and changes nothing until '
        'you confirm. Optionally, run Verify-Compatibility.cmd first: it checks whether '
        'your OpenCode build can be patched and changes nothing.')

    heading(document, 'Install', level=1)
    for step in [
        'Quit OpenCode completely (all windows and the tray icon).',
        'Double-click Install-Sidebar.cmd and wait for the message '
        '\u201cInstalled. Original archive is backed up; launch OpenCode normally.\u201d',
        'Start OpenCode and open any session.',
    ]:
        document.add_paragraph(step, style='List Number')
    document.add_paragraph(
        'You should see a Telemetry button at the bottom-right of the window, with a Search '
        'button beside it. The sidebar is open by default; click the button (or press Escape '
        'while it has focus) to collapse it. The first refresh can take a few seconds while '
        'the session loads.')

    heading(document, 'What it looks like when it works', level=1)
    document.add_paragraph(
        'The panel below was captured on a real session. Compare it with what you see: '
        'the same sections, the same labels, live numbers from your own account.')
    for name, caption, explanation in SCREENSHOTS:
        heading(document, caption, level=2)
        document.add_paragraph(explanation)
        document.add_picture(str(MEDIA / name), width=Inches(6))
        last = document.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = document.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.font.size = Pt(9)
        run.italic = True

    heading(document, 'Search', level=1)
    document.add_paragraph(
        'The Search button next to Telemetry searches the text of every session stored by '
        'OpenCode \u2014 user prompts and agent replies \u2014 and opens any hit as an additional '
        'session tab, leaving your open sessions untouched.')
    bullets(document, [
        'The User and Agent checkboxes choose which roles are searched (both on by default).',
        'The dropdown caps the number of results; an empty query lists the most recent messages.',
        'Hover a result for a wider context window; reasoning and tool calls are never searched.',
        'Text is escaped before rendering, the database is only ever read, and a search that '
        'finds nothing does not slow the app down: typing again cancels the previous search.',
    ])

    heading(document, 'Reading the numbers honestly', level=1)
    bullets(document, [
        'Window totals (1% used, 64% used, \u2026) come from the provider and are exact.',
        'Per-model shares under the Go caps are estimates: the provider reports only a '
        'window total, so the sidebar distributes it by each model\u2019s share of locally '
        'recorded cost. Usage from other machines is not counted.',
        'A cap the provider does not report shows Unavailable, never a guessed number.',
        'Free models show 0% because they carry no billable cost.',
        'A Stale label means a refresh failed and the last good reading is kept on screen.',
    ])

    heading(document, 'Uninstall', level=1)
    for step in [
        'Quit OpenCode completely.',
        'Double-click Rollback-Sidebar.cmd.',
    ]:
        document.add_paragraph(step, style='List Number')
    document.add_paragraph(
        'The original app archive is restored and the telemetry plugin is moved into the '
        'backups folder rather than deleted, so nothing is lost.')

    heading(document, 'If something goes wrong', level=1)
    bullets(document, [
        'Installer says OpenCode is still running: quit it fully, including the tray icon.',
        'Installer says the archive changed: OpenCode was updated after staging; just run '
        'Install-Sidebar.cmd again, it re-stages automatically.',
        'Sidebar missing after an OpenCode update: expected; run Install-Sidebar.cmd again.',
        'Everything shows Unavailable: check that you are signed in, then look at the small '
        'diagnostics line at the bottom of the panel (session id and observed call counts) '
        'and include it when asking for help.',
    ])

    heading(document, 'Safety notes', level=1)
    bullets(document, [
        'The original app archive is backed up before anything is replaced.',
        'Your credentials never leave the OpenCode server process; the sidebar only ever '
        'receives percentages and reset times.',
        'Nothing is uploaded anywhere. The only network requests are to the two provider '
        'usage endpoints.',
    ])

    document.save(OUT)
    print(f'wrote {OUT} ({OUT.stat().st_size / 1024:.0f} KiB)')


if __name__ == '__main__':
    main()
